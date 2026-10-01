import os
import json
import time
import requests
from ddgs import DDGS
from dotenv import load_dotenv
# from google import genai
from openpyxl import Workbook, load_workbook


# ============================================================
# CONFIGURATION
# ============================================================

SEARCH_URL = "http://localhost:8080/search"
QWEN_URL = "http://localhost:11434/api/generate"
QWEN_MODEL = "qwen3:1.7b"

# GEMINI_MODEL = "gemini-3.8-flash"

EXCEL_PATH = "pm_leads.xlsx"

COLUMNS = ["Company", "Type", "Score", "Reason", "Source URL"]

VELA_CONTEXT_PATH = "vela_context.md"


# ============================================================
# GEMINI
# ============================================================

# load_dotenv()

# GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")

# if not GEMINI_API_KEY:
#     raise RuntimeError("GEMINI_API_KEY is not set.")

# gemini_client = genai.Client(api_key=GEMINI_API_KEY)


# def call_gemini(prompt: str, model: str = GEMINI_MODEL, max_retries: int = 5) -> str:
#     """Call Gemini with automatic retry on 503 (model busy) errors."""
#     for attempt in range(max_retries):
#         try:
#             response = gemini_client.models.generate_content(
#                 model=model,
#                 contents=prompt
#             )
#             return response.text
#         except Exception as e:
#             if "503" in str(e) and attempt < max_retries - 1:
#                 wait = 5 * (attempt + 1)
#                 print(f"    Critic model busy, retrying in {wait}s...")
#                 time.sleep(wait)
#             else:
#                 raise


# def critic_review(analysis: dict, company_info: str) -> dict:
#     """Re-check every signal the Analyst marked True against the evidence."""
#     signals = analysis.get("signals", {})

#     true_signals = {
#         name: data.get("evidence", "")
#         for name, data in signals.items()
#         if data.get("supported") is True
#     }

#     if not true_signals:
#         return analysis

#     signals_text = "\n".join(
#         f"- {name}: claimed evidence = \"{evidence}\""
#         for name, evidence in true_signals.items()
#     )

#     prompt = f"""
# You are auditing another AI's evidence claims. Be skeptical, not agreeable.

# SEARCH RESULT:
# {company_info}

# The other AI marked these signals as TRUE and cited this evidence for each:

# {signals_text}

# RULE: A signal may only stay TRUE if the cited evidence is DIRECTLY stated
# in the search result above — not implied, assumed, or inferred from industry
# stereotypes. If the evidence is missing, vague, or not actually in the text,
# the signal must be marked FALSE.

# Return ONLY valid JSON in this exact structure, one entry per signal listed above:

# {{
#     "signal_name": {{"confirmed": true or false, "note": "short reason"}}
# }}
# """

#     answer = call_gemini(prompt)

#     try:
#         cleaned = answer.strip().removeprefix("```json").removesuffix("```").strip()
#         critic_result = json.loads(cleaned)
#     except json.JSONDecodeError:
#         print("  Critic returned invalid JSON — keeping Analyst's original signals")
#         return analysis

#     for signal_name, verdict in critic_result.items():
#         if signal_name in signals and verdict.get("confirmed") is False:
#             print(f"  Critic overturned: {signal_name} — {verdict.get('note', '')}")
#             signals[signal_name]["supported"] = False

#     return analysis

# ============================================================
# VELA CONTEXT (NEW)
# ============================================================

def load_vela_context(path: str = VELA_CONTEXT_PATH) -> str:
    """Load Vela's business description from an external markdown file.

    Args:
        path: Path to the markdown file describing Vela's services and context.

    Returns:
        The file's raw text content, to be dropped straight into the prompt.
    """
    with open(path, "r", encoding="utf-8") as f:
        return f.read()


VELA_CONTEXT = load_vela_context()


# ============================================================
# SCORING RUBRIC
# ============================================================

SCORE_WEIGHTS = {
    "contract_opportunity": 25,
    "project_product_management_need": 20,
    "remote_or_hybrid": 15,
    "porto_gaia": 15,
    "tech_or_media_company": 10,
    "current_open_opportunity": 10,
    "decision_maker": 5,
}


# ============================================================
# EXCEL SAVING (NEW)
# ============================================================

def save_lead(company: str, lead_type: str, score: int, reason: str, url: str) -> bool:
    """Append a lead to pm_leads.xlsx, skipping it if the company is already saved.

    Args:
        company: The company name identified by the Analyst.
        lead_type: "opportunity", "potential_opportunity", or "irrelevant".
        score: The calculated numeric score (0-100).
        reason: The Analyst's short explanation for the classification.
        url: The source URL the lead was found at.

    Returns:
        True if the lead was newly added, False if it was skipped as a duplicate.
    """
    # Open the existing file, or start a new one with headers if it doesn't exist yet
    if os.path.exists(EXCEL_PATH):
        workbook = load_workbook(EXCEL_PATH)
        sheet = workbook.active
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(COLUMNS)

    # Check every existing row's company name (column A) for a case-insensitive match
    existing_companies = {
        str(row[0].value).strip().lower()
        for row in sheet.iter_rows(min_row=2)
        if row[0].value
    }

    if company.strip().lower() in existing_companies:
        print(f"  Skipped (already saved): {company}")
        return False

    sheet.append([company, lead_type, score, reason, url])
    workbook.save(EXCEL_PATH)
    print(f"  Saved: {company}")
    return True


def run_query(query: str) -> None:
    """Search for one query phrase and analyze/save each result.

    Args:
        query: The search phrase to send to SearXNG.
    """
    # ============================================================
    # SEARCH
    # ============================================================

    search_response = requests.get(
        SEARCH_URL,
        params={
            "q": query,
            "format": "json"
        },
        timeout=30
    )

    search_response.raise_for_status()

    results = search_response.json().get("results", [])[:5]

    print(f"  SearXNG returned {len(results)} results")

    if not results:
        print("  SearXNG returned no usable results for this query.")
        return

    # Remove duplicate URLs from this search
    seen_urls = set()
    unique_results = []

    for result in results:
        url = result.get("url", "").strip()

        if url and url not in seen_urls:
            seen_urls.add(url)
            unique_results.append(result)

    results = unique_results

    results = [result for result in results if result.get("url", "").strip()]

    # ============================================================
    # PROCESS EACH SEARCH RESULT
    # ============================================================

    for result in results:

        company_info = f"""
    Title: {result.get("title", "")}
    Description: {result.get("content", "")}
    URL: {result.get("url", "")}
    """

        prompt = f"""
You are qualifying business opportunities for Vela Strategies.

{VELA_CONTEXT}

Analyze ONLY the search result provided below.

SEARCH RESULT:
{company_info}

Your job is to determine whether this result identifies a useful business
opportunity for Vela.

A useful opportunity is evidence that an identifiable organization in Portugal
has a current or recent need for project management or product management
services that could potentially be provided by an external professional
through Vela.

============================================================
CLASSIFICATION
============================================================

Classify the result as exactly one of:

- "opportunity"
  Direct evidence that the organization has a current or recent
  project/product-management need AND the engagement may be suitable
  for an external provider.

  This includes:
  - contract work
  - freelance work
  - consulting
  - temporary work
  - fixed-term work
  - project-based work
  - fractional work
  - contractor arrangements
  - situations where the organization explicitly accepts contractors,
    consultants, agencies, or external service providers

- "potential_opportunity"
  The organization has a clear current or recent project/product-
  management need, but the result does NOT provide evidence that the
  organization is open to external, contract, freelance, consulting,
  temporary, or similar arrangements.

  A normal permanent/full-time Project Manager or Product Manager
  employee vacancy should generally be classified as
  "potential_opportunity".

- "irrelevant"
  The result is not useful for this campaign.

IMPORTANT:

A Project Manager or Product Manager job posting is NOT automatically
an "opportunity" for Vela.

A job posting can be an "opportunity" only when there is evidence that
the engagement could be performed by an external provider.

If there is a relevant PM/product-management vacancy but the engagement
type is unclear, classify it as "potential_opportunity".

Do not assume that a company accepts contractors merely because:
- the role is in technology
- the role is remote or hybrid
- the company is a startup
- the role is senior
- the role is project-based
- the company has used contractors in the past
- the job appears on a freelance website

Only mark an opportunity as suitable for external work when the result
contains direct evidence of that arrangement.

Job boards, recruiting platforms, freelancer marketplaces, and
directories MAY be used as sources.

A job-board result can be useful when it contains a relevant
project/product-management vacancy that identifies the actual
employer.

In that case:

- The job board is the SOURCE.
- The employer is the COMPANY.
- The employer should be recorded as the company.
- The job board URL may be retained as the source URL.

Do NOT record the job board itself as the company.

For example, if a DailyRemote page contains a Project Manager
vacancy for "Company X", identify Company X as the company.

If a job board page only contains a general collection of jobs and
does not identify a specific relevant employer, classify it as
"irrelevant".

A normal employee PM/Product Manager vacancy can still be a
"potential_opportunity". It is useful because the partner may
research the employer and determine whether they would consider
contract, freelance, consulting, or external-provider arrangements.

Do not assume that an employer is open to contract work merely
because the vacancy appears on a job board.

============================================================
SIGNALS
============================================================

Evaluate these EXACT signals.

1. contract_opportunity
+25 points

Is there direct evidence of contract, freelance, consulting, consultant,
fractional, temporary, fixed-term, part-time, project-based, or external
professional work?

2. project_product_management_need
+20 points

Is there direct evidence of a need for project management, product management,
technical project management, program management, delivery management,
product operations, or similar work?

3. remote_or_hybrid
+15 points

Is there direct evidence that the opportunity is remote or hybrid?

4. porto_gaia
+15 points

Is there direct evidence connecting the organization or opportunity to
Vila Nova de Gaia, Porto, or the Greater Porto area?

5. tech_or_media_company
+10 points

Is there direct evidence that the organization is a technology, software,
SaaS, startup, product, digital agency, technology consulting, digital media,
or technology/media company?

6. current_open_opportunity
+10 points

Is there direct evidence that the opportunity is currently open, active,
recently posted, or otherwise currently actionable?

7. decision_maker
+5 points

Is an identifiable relevant person mentioned, such as a founder, CEO, CTO,
COO, operations manager, product leader, project leader, or hiring manager?

============================================================
EVIDENCE RULE
============================================================

This is the most important rule.

A signal may ONLY be marked true when the SEARCH RESULT itself contains
direct evidence supporting that signal.

Do not use outside knowledge.

Do not guess.

Do not infer from industry stereotypes.

If evidence is missing or ambiguous, mark the signal false.

For every TRUE signal, provide the specific evidence from the search result.

For FALSE signals, leave the evidence field empty.

Examples:

- A company being located in Porto supports porto_gaia.
- A company being a technology company does NOT automatically support
  project_product_management_need.
- A company hiring does NOT automatically mean it needs project management.
- A remote job does NOT automatically mean it is a contract opportunity.
- A company's growth does NOT automatically mean it has a project-management need.
- A company being in the technology industry supports tech_or_media_company,
  but does not by itself prove there is a current opportunity.

============================================================
PORTUGAL REQUIREMENT
============================================================

Only consider opportunities located in Portugal.

If the search result clearly concerns an opportunity outside Portugal,
classify it as "irrelevant".

============================================================
OUTPUT
============================================================

Return ONLY valid JSON.

Use exactly this structure:

{{
    "company": "actual organization name",
    "type": "opportunity",
    "signals": {{
        "contract_opportunity": {{
            "supported": false,
            "evidence": ""
        }},
        "project_product_management_need": {{
            "supported": false,
            "evidence": ""
        }},
        "remote_or_hybrid": {{
            "supported": false,
            "evidence": ""
        }},
        "porto_gaia": {{
            "supported": false,
            "evidence": ""
        }},
        "tech_or_media_company": {{
            "supported": false,
            "evidence": ""
        }},
        "current_open_opportunity": {{
            "supported": false,
            "evidence": ""
        }},
        "decision_maker": {{
            "supported": false,
            "evidence": ""
        }}
    }},
    "reason": "short explanation of the classification"
}}

Do NOT calculate a score.
Python will calculate the score from the signals.
"""


        # ========================================================
        # ASK QWEN TO ANALYZE THE RESULT
        # ========================================================

        response = requests.post(
    QWEN_URL,
    json={
        "model": QWEN_MODEL,
        "prompt": prompt,
        "stream": False,
        "format": "json",
        "think": False
    },
    timeout=120
)

        response.raise_for_status()

        answer = response.json()["response"]


        # ========================================================
        # PARSE QWEN'S JSON
        # ========================================================

        try:
            analysis = json.loads(answer)

        except json.JSONDecodeError:

            print("\nQwen returned invalid JSON:")
            print(answer)
            print("-" * 60)

            continue

        # ========================================================
        # CRITIC REVIEW
        # ========================================================

        # analysis = critic_review(analysis, company_info)

        # ========================================================
        # CALCULATE SCORE
        # ========================================================

        score = 0

        signals = analysis.get("signals", {})

        for signal, points in SCORE_WEIGHTS.items():

            signal_data = signals.get(signal, {})

            if signal_data.get("supported") is True:
                score += points


        # ========================================================
        # DISPLAY RESULT
        # ========================================================

        print("\n" + "=" * 60)

        print(f"Company: {analysis.get('company', 'Unknown')}")
        print(f"Type: {analysis.get('type', 'Unknown')}")
        print(f"Score: {score}/100")

        print("\nSignals:")

        for signal, points in SCORE_WEIGHTS.items():

            signal_data = signals.get(signal, {})

            supported = signal_data.get("supported", False)
            evidence = signal_data.get("evidence", "")

            if supported:
                print(f"  [+{points}] {signal}")
                print(f"       Evidence: {evidence}")

            else:
                print(f"  [ 0] {signal}")

        print("\nReason:")
        print(analysis.get("reason", ""))

        print("\nSource:")
        print(result.get("url", ""))

        print("=" * 60)


        # ========================================================
        # SAVE TO EXCEL
        # ========================================================
        # Save both types of relevant opportunities.
        # Irrelevant results are discarded.

        if analysis.get("type") in ["opportunity", "potential_opportunity"]:
            save_lead(
                company=analysis.get("company", "Unknown"),
                lead_type=analysis.get("type", "Unknown"),
                score=score,
                reason=analysis.get("reason", ""),
                url=result.get("url", ""),
            )
        else:
            print(
                f"  Not saved ({analysis.get('type', 'unknown')}): "
                f"{analysis.get('company', 'Unknown')}"
            )

# ============================================================
# RUN ALL QUERIES
# ============================================================

QUERIES = [
    "Portugal project manager",
    "Portugal project management",
    "Portugal product manager",
    "Portugal product management",
    "Portugal technical project manager",
    "Portugal digital project manager",
    "Portugal project manager freelance",
    "Portugal project manager contract",
    "Portugal project manager consultant",
    "Portugal project manager remote",
]

for search_query in QUERIES:
    print(f"\n\n### Running query: {search_query} ###")
    run_query(search_query)