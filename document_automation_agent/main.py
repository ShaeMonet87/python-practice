import subprocess
from pathlib import Path

result = subprocess.run(
    [
        "osascript",
        "-e",
        'tell application "Notes" to get plaintext of note "Billable Hours"'
    ],
    capture_output=True,
    text=True
)

note_text = result.stdout

lines = note_text.splitlines()

print("Number of lines:", len(lines))

entries = []
current_entry = {}

for line in lines:
    if line.startswith("Date:"):
        if current_entry:
            entries.append(current_entry)

        current_entry = {
            "date": line.replace("Date:", "").strip()
        }

    elif line.startswith("Project:"):
        current_entry["project"] = line.replace("Project:", "").strip()

    elif line.startswith("Hours:"):
        current_entry["hours"] = line.replace("Hours:", "").strip()

    elif line.startswith("Description:"):
        current_entry["description"] = line.replace("Description:", "").strip()

if current_entry:
    entries.append(current_entry)

print("Entries found:", len(entries))


for entry in entries:
    if "hours" in entry:
        entry["hours"] = float(entry["hours"])

for i, entry in enumerate(entries, start=1):
    print(f"Entry {i}: {len(entry)} fields")

excel_file = Path(__file__).parent / "billable_hours.xlsx"

from openpyxl import Workbook, load_workbook

if excel_file.exists():
    workbook = load_workbook(excel_file)
    sheet = workbook["Billable Hours"]
else:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Billable Hours"
    sheet.append(["Date", "Project", "Hours", "Description"])

existing_entries = set()

for row in sheet.iter_rows(min_row=2, values_only=True):
    date = row[0]
    project = row[1]
    hours = row[2]
    description = row[3]

    existing_entries.add((date, project, hours, description))

for entry in entries:
    project = entry.get("project", "")
    hours = entry.get("hours", "")
    description = entry.get("description", "")

    entry_key = (
    entry.get("date", ""),
    project,
    hours,
    description
)

    if entry_key not in existing_entries:
        sheet.append([
            entry.get("date", ""),
            project,
            hours,
            description
        ])
        existing_entries.add(entry_key)

workbook.save(excel_file)

print("Excel file updated.")

